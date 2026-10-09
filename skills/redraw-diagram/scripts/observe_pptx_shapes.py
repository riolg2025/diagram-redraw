#!/usr/bin/env python3
"""
observe_pptx_shapes.py —— 入2：把 PPT 里的一页**清点**成 observations.json。

★ 它产出的是「观察」，不是「清单」。

    观察 = 原件里有什么   —— 机械清点，可复现，**只增不改**（证据）
    清单 = 这张图在说什么 —— 要判断、要归并、要问用户（结论）

    "有 18 个形状" 是观察；"这是一条流程、有 10 条关系" 是结论。
    两者不能混 —— 混了就是拿**画法**属性冒充**意义**属性。
    项目栽过：11 个线条对象，10 条关系，**两个数都对**。

所以这个脚本**不猜**：形状就是形状，线就是线，连接符的端点有就是有、没有就是没有。
归并成"一条关系"是下一步（agent）的事。

用法：
    python3 observe_pptx_shapes.py deck.pptx --slide 6 --out work/observations.json
    python3 observe_pptx_shapes.py deck.pptx --slide 5 --region 747,209,1224,616 \
            --render work/rendered/page-0005.png --out work/observations.json

--region 是**原页面像素坐标**，只用来**筛**（哪些对象属于这一块），**不做坐标变换**。
--render 是**整页**渲染图，用来采"继承主题、XML 里没写"的那些颜色。
"""

import argparse
import json
import os
import sys
from collections import Counter

from PIL import Image
from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from common import Xfrm, GroupCtx, emu_to_px, mode_fill, darkest, lum  # noqa: E402

A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
P = '{http://schemas.openxmlformats.org/presentationml/2006/main}'

OBS_VERSION = "0.1"

ARROW_PRSTS = {'upArrow', 'downArrow', 'leftArrow', 'rightArrow', 'upDownArrow',
               'leftRightArrow', 'bentArrow', 'bentUpArrow', 'curvedDownArrow',
               'curvedUpArrow', 'curvedLeftArrow', 'curvedRightArrow',
               'stripedRightArrow', 'notchedRightArrow', 'circularArrow',
               'quadArrow', 'leftRightUpArrow'}
BRACE_PRSTS = {'leftBrace', 'rightBrace'}


# --------------------------------------------------------------------------
# XML 小工具
# --------------------------------------------------------------------------
def _first(el, tag):
    for e in el.iter():
        if e.tag == tag:
            return e
    return None


def theme_colors(pptx_path):
    """读原件的主题色表：{accent1: "08266E", ...}。

    ★ 为什么必须解析：搬运过来的形状用的是 `schemeClr`（**主题色**）。
      不解析的话，它会在**新文档的主题**下重新取色 —— 实测 8 个图标
      从深藏青变成了浅蓝。**主题色是原件的事实，不能让它跟着新主题漂。**
    """
    import zipfile
    out = {}
    try:
        z = zipfile.ZipFile(pptx_path)
        name = next((n for n in z.namelist()
                     if n.startswith("ppt/theme/theme") and n.endswith(".xml")), None)
        if not name:
            return out
        root = etree.fromstring(z.read(name))
        cs = next((e for e in root.iter()
                   if e.tag.rsplit('}', 1)[-1] == 'clrScheme'), None)
        if cs is None:
            return out
        for e in cs:
            key = e.tag.rsplit('}', 1)[-1]
            v = next((c.get('val') for c in e.iter() if c.get('val')), None)
            if v:
                out[key] = v
        # ★ bg1 / tx1 / bg2 / tx2 **不在 clrScheme 里** —— 它们是**别名**，
        #   映射表在母版的 `p:clrMap` 上（bg1→lt1 之类）。
        #   不做这一步，`schemeClr val="bg1"` 就解析不出来。
        #   （实测：不做的话 32 处会掉进"找不到"，靠新加的告警才发现。）
        for n in z.namelist():
            if not (n.startswith("ppt/slideMasters/slideMaster")
                    and n.endswith(".xml")):
                continue
            root = etree.fromstring(z.read(n))
            cm = next((e for e in root.iter()
                       if e.tag.rsplit('}', 1)[-1] == 'clrMap'), None)
            if cm is not None:
                for k, v in cm.attrib.items():
                    if k not in out and v in out:
                        out[k] = out[v]
            break
    except Exception:
        pass
    return out


def _rgb_to_hsl(r, g, b):
    r, g, b = r / 255.0, g / 255.0, b / 255.0
    mx, mn = max(r, g, b), min(r, g, b)
    l = (mx + mn) / 2
    if mx == mn:
        return 0.0, l, 0.0
    d = mx - mn
    s = d / (2 - mx - mn) if l > 0.5 else d / (mx + mn)
    if mx == r:
        h = ((g - b) / d) % 6
    elif mx == g:
        h = (b - r) / d + 2
    else:
        h = (r - g) / d + 4
    return h / 6.0, l, s


def _hsl_to_rgb(h, l, s):
    if s == 0:
        v = int(round(l * 255))
        return v, v, v

    def f(p, q, t):
        t = t % 1.0
        if t < 1 / 6.0:
            return p + (q - p) * 6 * t
        if t < 1 / 2.0:
            return q
        if t < 2 / 3.0:
            return p + (q - p) * (2 / 3.0 - t) * 6
        return p
    q = l * (1 + s) if l < 0.5 else l + s - l * s
    p = 2 * l - q
    return tuple(int(round(max(0.0, min(1.0, f(p, q, h + o))) * 255))
                 for o in (1 / 3.0, 0.0, -1 / 3.0))


def apply_color_mods(hexv, el):
    """把 DrawingML 的颜色修饰符（lumMod / lumOff / tint / shade）算进颜色。

    ★ 为什么不能只查表了事：主题色常常带修饰符（比如 accent1 调暗 60%）。
      只取基色会**悄悄给一个错的颜色**。算不出来的修饰符会**报出来**，不静默。
    """
    if hexv is None:
        return None, []
    r, g, b = (int(hexv[i:i + 2], 16) for i in (0, 2, 4))
    h, l, sat = _rgb_to_hsl(r, g, b)
    unknown = []
    for c in el:
        tag = c.tag.rsplit('}', 1)[-1]
        try:
            v = int(c.get('val')) / 100000.0
        except Exception:
            continue
        if tag == 'lumMod':
            l *= v
        elif tag == 'lumOff':
            l += v
        elif tag == 'tint':
            l = l * (1 - v) + v
        elif tag == 'shade':
            l *= v
        elif tag in ('alpha', 'alphaMod', 'alphaOff'):
            pass                      # 透明度不改颜色
        elif tag in ('satMod', 'satOff', 'hueMod', 'hueOff', 'comp', 'inv',
                     'gamma', 'invGamma'):
            unknown.append(tag)
    l = max(0.0, min(1.0, l))
    return "%02X%02X%02X" % _hsl_to_rgb(h, l, sat), unknown


def resolve_theme_colors(el, scheme):
    """把 XML 里的 `schemeClr` / `sysClr` 就地换成解析后的 `srgbClr`。

    ⚠️ 只处理**没有修饰符**的（lumMod / tint / shade …）。
      有修饰符就**返回警告**，不静默给一个错的颜色 ——
      这个项目栽过太多次"悄悄错了"。
    """
    warn = []
    for e in list(el.iter()):
        ln = e.tag.rsplit('}', 1)[-1]
        if ln not in ('schemeClr', 'sysClr'):
            continue
        key = e.get('val')
        hexv = scheme.get(key) or (e.get('lastClr') if ln == 'sysClr' else None)
        mods = [c.tag.rsplit('}', 1)[-1] for c in e]
        if mods:
            warn.append("%s=%s 带修饰符 %s —— 只换了底色，**亮度没算**，可能不准"
                        % (ln, key, ",".join(mods)))
        if not hexv:
            warn.append("%s=%s 在原件主题里找不到，保持原样（新主题下会变色）"
                        % (ln, key))
            continue
        parent = e.getparent()
        idx = list(parent).index(e)
        new = etree.Element(A + 'srgbClr')
        new.set('val', hexv)
        for c in list(e):
            new.append(c)
        parent.remove(e)
        parent.insert(idx, new)
    return warn


def _first_local(el, local):
    """按**元素名**找，不管命名空间。

    ★ 为什么不能写死命名空间：`<p:spPr>` 在 **presentationml** 里，
      而它的孩子 `<a:solidFill>` 在 **drawingml** 里。
      写成 `a:spPr` 就永远找不到 —— 后果是**所有形状的填充都读不到**，
      重画出来底色全丢，而且不报错。
    """
    for e in el.iter():
        if e.tag.rsplit('}', 1)[-1] == local:
            return e
    return None


def prst_of(el):
    g = _first(el, A + 'prstGeom')
    return (g.get('prst') or '') if g is not None else ''


def text_of(sh):
    try:
        return sh.text_frame.text.strip() if sh.has_text_frame else ""
    except Exception:
        return ""


def sp_pr_of(el):
    """p:sp / p:cxnSp 的 <p:nvSpPr><p:cNvPr>，里面有 id 和 name。"""
    for tag in ('cNvPr',):
        for e in el.iter():
            if e.tag.endswith('}' + tag) and e.get('id'):
                return e
    return None


def cnv_id(el):
    e = sp_pr_of(el)
    return int(e.get('id')) if e is not None else None


def cnv_name(el):
    e = sp_pr_of(el)
    return e.get('name') if e is not None else None


def ln_props(el, scheme=None):
    """线条：宽度、虚线、**XML 里的显式颜色**（没写就是 None，表示继承主题）。

    还读两样实测踩过的东西：
      alpha —— 原图是 91% 不透明度，不读就写成 100%，线会更"实"
      arrow —— 箭头**形状**（arrow / triangle / stealth…）。实测原图是 "arrow"
               （细长），硬编码成 "triangle"（实心三角）后头部明显大一圈。
    """
    ln = _first(el, A + 'ln')
    base = {"pt": 0.75, "dash": "solid", "color": None, "src": "default",
            "alpha": None, "arrow": None}
    if ln is None:
        return base
    w = int(ln.get('w')) if ln.get('w') else 9525
    pd = ln.find(A + 'prstDash')
    dash = pd.get('val') if pd is not None else 'solid'
    te = ln.find(A + 'tailEnd')
    he = ln.find(A + 'headEnd')
    arrow = None
    if te is not None and te.get('type') not in (None, 'none'):
        arrow = "both" if (he is not None
                           and he.get('type') not in (None, 'none')) else "end"
    base.update({"pt": round(w / 12700.0, 3), "dash": dash, "arrow": arrow})
    if te is not None and te.get('type'):
        base["arrow_type"] = te.get('type')
    if ln.find(A + 'noFill') is not None:
        base["src"] = "nofill"
        return base
    sf = ln.find(A + 'solidFill')
    if sf is not None:
        c = sf.find(A + 'srgbClr')
        if c is not None:
            base["color"] = c.get('val')
            base["src"] = "xml"
            al = c.find(A + 'alpha')
            if al is not None and al.get('val'):
                base["alpha"] = int(al.get('val'))
            return base
        sc = sf.find(A + 'schemeClr')
        if sc is not None:
            # ★ 主题色描边要**解析**它，既不许丢、也不许采。
            #   原来这里 color 留 None，而下面的采样条件只认 default/unset，
            #   于是 schemeClr 的边框被**静默丢掉**（实测造用例复现过）。
            hexv, unk = apply_color_mods((scheme or {}).get(sc.get('val')), sc)
            base["src"] = "scheme" if hexv else "scheme_unresolved"
            base["scheme"] = sc.get('val')
            base["unknown_mods"] = unk
            if hexv:
                base["color"] = hexv
            return base
    base["src"] = "unset"
    return base


def fill_of(el, scheme=None):
    """填充。**每一种情况都必须是不同的状态** —— 不能都落到 color=None。

    ★ 这里原来把三件不同的事挤进了同一个 color=None：
        · 真的没写填充        （unset）
        · 渐变 / 图案 / 图片填充（也是 unset！）→ 会被"采成一个平色"，等于把
          渐变**假装成**纯色画出去
        · 主题色              （scheme）      → 不去解析，反而去采样
      实测：造一页 gradFill，观察器吐出来和"没写填充"一模一样。
    """
    sp = _first_local(el, 'spPr')      # ★ p:spPr，不是 a:spPr
    if sp is None:
        return {"color": None, "src": "nospPr"}
    if sp.find(A + 'noFill') is not None:
        return {"color": None, "src": "nofill"}
    # ★ 渐变 / 图案 / 图片：**单独的状态**，不许和"没写"混在一起
    for tag in ('gradFill', 'pattFill', 'blipFill', 'grpFill'):
        if sp.find(A + tag) is not None:
            return {"color": None, "src": tag.replace('Fill', ''), "what": tag}
    sf = sp.find(A + 'solidFill')
    if sf is None:
        # 没写 fill 元素：主题给了就是主题的，没给就是**没有填充**
        return {"color": None,
                "src": "themefill" if has_theme_fill(el) else "nofill"}
    c = sf.find(A + 'srgbClr')
    if c is not None:
        return {"color": c.get('val'), "src": "xml"}
    s = sf.find(A + 'schemeClr')
    if s is not None:
        # ★ 主题色是**明写的**，只是要查主题表 —— 和文字色同一条规矩。
        hexv, unk = apply_color_mods((scheme or {}).get(s.get('val')), s)
        return {"color": hexv, "src": "scheme" if hexv else "scheme_unresolved",
                "scheme": s.get('val'), "unknown_mods": unk}
    return {"color": None, "src": "other"}


def font_of(el, scheme=None):
    """字号 / 字体。

    ★ 原来只读**第一个 `<a:r>` 的 rPr`**。可是字号常常写在 `a:defRPr` 里
      （段落默认），这样就读成 None —— 而下游把 None 静默变成 10pt / 微软雅黑。
      实测：造一个只在 defRPr 里写 sz=9900 的形状，观察器给出 {pt:None, face:None}。

      和 text_color_of 保持一致：**rPr 和 defRPr 都看**。
      两个都读不到就如实标 unknown，让下游去报，**不许静默填默认值**。
    """
    tb = _first(el, P + 'txBody')
    if tb is None:
        return {"pt": None, "face": None, "src": "none"}
    pt = face = None
    for tag in ('rPr', 'defRPr'):
        for rPr in tb.iter(A + tag):
            if pt is None and rPr.get('sz'):
                pt = int(rPr.get('sz')) / 100.0
            if face is None:
                for t in ('latin', 'ea'):
                    e = rPr.find(A + t)
                    if e is not None and e.get('typeface'):
                        face = e.get('typeface')
                        break
    if pt is None and face is None:
        return {"pt": None, "face": None, "src": "unknown"}
    return {"pt": pt, "face": face, "src": "xml"}


def text_color_of(el, scheme=None):
    """文字颜色，**先从 XML 读**。

    ★ 规矩和线条一样：**XML 里显式写了就用它，采样只是渲染的近似。**
      漏了这条的后果实测过：蓝底白字的框，用"框内最深像素"去采，
      采到的是**蓝底本身** → 文字被写成背景色 → **整段字看不见**。
    """
    tb = _first(el, P + 'txBody')
    if tb is None:
        return {"color": None, "src": "none"}
    for tag in ('rPr', 'defRPr'):
        for rPr in tb.iter(A + tag):
            sf = rPr.find(A + 'solidFill')
            if sf is None:
                continue
            c = sf.find(A + 'srgbClr')
            if c is not None:
                return {"color": c.get('val'), "src": "xml"}
            s = sf.find(A + 'schemeClr')
            if s is not None:
                key = s.get('val')
                # ★ schemeClr 是**显式写的**，只是要拿主题表解析 —— 不能跑去采样。
                #   实测：绿条上的白字标题（schemeClr lt1）被采成了深藏青。
                base = (scheme or {}).get(key)
                hexv, unk = apply_color_mods(base, s)
                return {"color": hexv, "src": "scheme" if hexv else "scheme_unresolved",
                        "scheme": key, "unknown_mods": unk}
    return {"color": None, "src": "unset"}


def text_color_sample(img, box, fill_hex, tol=110):
    """兜底：从渲染图里找文字色。

    不能取"最深像素" —— 白字在深底上会被采成底色。
    改成：**框内出现最多的那个"和填充色明显不同"的颜色**。
    文字像素数仅次于底色，所以这条稳。
    """
    from collections import Counter
    x0, y0 = int(box[0]), int(box[1])
    x1, y1 = int(box[0] + box[2]), int(box[1] + box[3])
    cnt = Counter()
    for yy in range(y0 + 2, y1 - 2):
        for xx in range(x0 + 2, x1 - 2):
            if 0 <= xx < img.width and 0 <= yy < img.height:
                cnt[img.getpixel((xx, yy))] += 1
    if not cnt:
        return None
    # ★ 拿不到填充色时（文字框通常是 noFill），**不能用"取最深像素"** ——
    #   实测踩过：绿条上的白字标题被采成了黑色。
    #   改成用**框内众数**当背景参考：文字框里众数就是它背后的底色。
    if fill_hex:
        ref = tuple(int(fill_hex[i:i + 2], 16) for i in (0, 2, 4))
    else:
        ref = cnt.most_common(1)[0][0]
    for px, _n in cnt.most_common(60):
        if sum(abs(a - b) for a, b in zip(px, ref)) > tol:
            return "%02X%02X%02X" % px
    return None


def has_theme_fill(el):
    """主题有没有给这个形状指定填充（`p:style/a:fillRef idx != 0`）。

    和描边是同一个分水岭：没写 fill 元素 ≠ 没有填充，
    还要看主题给没给。
    """
    for st in el.iter(P + 'style'):
        for r in st.iter(A + 'fillRef'):
            if (r.get('idx') or "0") not in ("0", ""):
                return True
    return False


def has_theme_line(el):
    """主题有没有给这个形状指定描边（`p:style/a:lnRef idx != 0`）。

    ★ 这是"没写 `<a:ln>`"和"没有边框"之间的分水岭：
        · 没有 `<a:ln>`、也没有 lnRef  → **原件就是没有边框**
        · 没有 `<a:ln>`、但有 lnRef    → 主题给了边框（要去渲染图看它长什么样）
      实测栽过：把前一种也当成"不知道"，去渲染图"找最深的像素"，
      结果把**文字自己**（或图形自己的填充）当成了边框色，
      凭空给 4 个形状加上了深色边框。
    """
    for st in el.iter(P + 'style'):
        for r in st.iter(A + 'lnRef'):
            if (r.get('idx') or "0") not in ("0", ""):
                return True
    return False


def para_of(el):
    """段落级排版：行距 lnSpc、段前 / 段后距。

    ★ 实测漏了它的后果：原图文字卡是 lnSpc=120000（120% 行距），
      不读就按默认 100% 渲染 —— **整块文字的位置就不一样了**。
    """
    tb = _first(el, P + 'txBody')
    if tb is None:
        return {}
    out = {}
    for tag, key in (('lnSpc', 'lnspc'), ('spcBef', 'spc_before'),
                     ('spcAft', 'spc_after')):
        for p in tb.iter(A + 'p'):
            pPr = p.find(A + 'pPr')
            if pPr is None:
                continue
            e = pPr.find(A + tag)
            if e is None or len(e) == 0:
                continue
            v = e[0]
            if v.get('val') is not None:
                out[key] = int(v.get('val'))
            elif v.get('pts') is not None:
                out[key] = int(v.get('pts'))
            break
        else:
            continue
    return out


def insets_of(el):
    tb = _first(el, P + 'txBody')
    bp = tb.find(A + 'bodyPr') if tb is not None else None
    if bp is None:
        return None
    def g(k, d):
        return emu_to_px(int(bp.get(k))) if bp.get(k) else d
    return {"l": g('lIns', 15.0), "t": g('tIns', 7.5),
            "r": g('rIns', 15.0), "b": g('bIns', 7.5),
            "anchor": bp.get('anchor') or 't',
            # ★ wrap 和 overflow 是两件事，别混
            "wrap": bp.get('wrap') or 'square'}


def cxn_ends(el):
    """连接符的端点绑定：<a:stCxn id= idx=/> / <a:endCxn .../>。没有就是 None。"""
    out = {}
    for tag, key in (('stCxn', 'st'), ('endCxn', 'end')):
        e = _first(el, A + tag)
        if e is not None and e.get('id'):
            out[key] = {"shape": int(e.get('id')),
                        "idx": int(e.get('idx')) if e.get('idx') else None}
        else:
            out[key] = None
    return out


def line_pts(box, xf):
    """一条线的两个端点（考虑 flip）。"""
    x0, y0, w, h = box
    x1, y1, x2, y2 = x0, y0, x0 + w, y0 + h
    if xf.flipH:
        x1, x2 = x2, x1
    if xf.flipV:
        y1, y2 = y2, y1
    return [x1, y1], [x2, y2]


# --------------------------------------------------------------------------
def in_region(box, region):
    if region is None:
        return True
    rx, ry, rw, rh = region
    cx, cy = box[0] + box[2] / 2, box[1] + box[3] / 2
    return rx <= cx <= rx + rw and ry <= cy <= ry + rh


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("pptx")
    ap.add_argument("--slide", type=int, required=True)
    ap.add_argument("--render", help="**整页**渲染图，用来采主题色")
    ap.add_argument("--region", help="只要页面上的一块：x,y,w,h（原页面像素，只筛不变换）")
    ap.add_argument("--out", default="observations.json")
    ap.add_argument("--media-dir", help="图片导出目录（默认 <out>.media）")
    args = ap.parse_args()

    region = [float(v) for v in args.region.split(",")] if args.region else None
    ren = Image.open(args.render).convert("RGB") if args.render else None
    if ren is None:
        print("⚠ 没给 --render：XML 里没写颜色（继承主题）的对象会缺颜色。", file=sys.stderr)

    prs = Presentation(args.pptx)
    slide = prs.slides[args.slide - 1]
    canvas = [emu_to_px(prs.slide_width), emu_to_px(prs.slide_height)]
    scheme = theme_colors(args.pptx)
    color_warnings = []
    degrade = []          # 表达不了、**不许假装**的那些（渐变填充 / 自由曲线…）

    media_dir = args.media_dir or (os.path.splitext(os.path.abspath(args.out))[0] + ".media")
    os.makedirs(media_dir, exist_ok=True)
    # ★ 「原样搬运」的 sidecar：表达不了的形状，把它的 XML 存下来
    verbatim_dir = os.path.splitext(os.path.abspath(args.out))[0] + ".verbatim"
    os.makedirs(verbatim_dir, exist_ok=True)
    out_dir = os.path.dirname(os.path.abspath(args.out))

    shapes, connectors, groups, unsupported = [], [], [], []
    z = 0
    pic_i = 0

    def walk(shapes_iter, ctx, group_path):
        nonlocal z, pic_i
        for sh in shapes_iter:
            el = sh._element
            xf = Xfrm(el, A)
            if not xf.ok:
                continue
            box = ctx.box_px(xf)
            sid = cnv_id(el)
            oid = "sp%d" % sid if sid is not None else "sp?"
            ev = ["p:%s#%s" % ("cxnSp" if el.tag == P + 'cxnSp' else "sp", sid)]

            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                if in_region(box, region):
                    groups.append({
                        "oid": oid, "shape_id": sid, "name": cnv_name(el),
                        "box": [round(v, 1) for v in box],
                        "rot": round(xf.rot / 60000.0),
                        "members": [("sp%d" % cnv_id(c._element))
                                    for c in sh.shapes if cnv_id(c._element) is not None],
                        "evidence": ev,
                    })
                walk(sh.shapes, ctx.child(xf), group_path + [oid])
                continue

            if not in_region(box, region):
                continue

            base = {
                "oid": oid, "shape_id": sid, "name": cnv_name(el),
                "type": str(sh.shape_type), "prst": prst_of(el),
                "box": [round(v, 1) for v in box],
                "rot": round(xf.rot / 60000.0),
                "flipH": xf.flipH, "flipV": xf.flipV,
                "group_path": group_path,
                "z": z,
                "text": text_of(sh),
                "evidence": ev,
            }
            z += 1

            # ---- 连接符：单独一张表 ----
            if el.tag == P + 'cxnSp' or str(sh.shape_type).startswith("LINE"):
                p1, p2 = line_pts(box, xf)
                ends = cxn_ends(el)
                connectors.append(dict(base, kind="connector",
                                       p1=[round(v, 1) for v in p1],
                                       p2=[round(v, 1) for v in p2],
                                       st_cxn=ends["st"], end_cxn=ends["end"],
                                       line=ln_props(el)))
                continue

            # ---- 图片 ----
            if str(sh.shape_type).startswith("PICTURE"):
                try:
                    ext = sh.image.ext
                    fn = os.path.join(media_dir, "img%02d.%s" % (pic_i, ext))
                    with open(fn, "wb") as fh:
                        fh.write(sh.image.blob)
                    src = os.path.relpath(fn, out_dir)
                except Exception as e:
                    src = None
                    unsupported.append(dict(base, what="图片导出失败", why=repr(e)))
                pic_i += 1
                shapes.append(dict(base, kind="picture", src=src))
                continue

            # ---- 表格 / 自由曲线：画不出来，如实记 ----
            if str(sh.shape_type).startswith("TABLE"):
                unsupported.append(dict(base, what="表格", why="不支持"))
                continue
            if str(sh.shape_type).startswith("FREEFORM"):
                # ★ 表达不了 ≠ 没法要。
                #   FREEFORM 是 <p:sp> + <a:custGeom>，**原样搬过来还是原生可编辑形状**。
                #   我一开始把它塞进"画不出来"，那是拿自己的管线能力当原件的边界。
                # ★ 主题色要按**原件的**主题解析成绝对色再搬，
                #   否则它会跟着新文档的主题漂（实测图标从藏青变浅蓝）。
                copy_el = etree.fromstring(etree.tostring(el))
                ws = resolve_theme_colors(copy_el, scheme)
                # 注意用 extend 不用 +=：嵌套函数里 += 会被当成局部变量
                color_warnings.extend("%s: %s" % (oid, w) for w in ws)
                fn = os.path.join(verbatim_dir, "%s.xml" % oid)
                with open(fn, "wb") as fh:
                    fh.write(etree.tostring(copy_el, xml_declaration=True,
                                            encoding="UTF-8", standalone=True))
                shapes.append(dict(base, kind="verbatim",
                                   xml=os.path.relpath(fn, out_dir)))
                continue

            # ---- 其余：形状 ----
            fl = fill_of(el, scheme)
            if fl["src"] in ("grad", "patt", "blip", "grp"):
                # ★ 渐变 / 图案 / 图片填充**画不了** —— 不许采成一个平色假装是纯色。
                #   原来它和"没写填充"共用 unset，会被采成平色：**把渐变假装成纯色**。
                degrade.append("%s: 用了 %s，第一版画不了（不假装成纯色）"
                               % (oid, fl.get("what") or fl["src"]))
            elif fl["color"] is None and fl["src"] in ("unset", "themefill",
                                                       "other") and ren is not None:
                # 这几种才该去渲染图采：主题给了但没写死 / 写了但读不出的形式
                mf = mode_fill(ren, box)
                if mf:
                    fl = {"color": mf, "src": "sampled"}
            if fl["src"] == "scheme_unresolved":
                color_warnings.append("%s: 填充的主题色 %s 在原件主题里找不到，已采样兜底"
                                      % (oid, fl.get("scheme")))
                if ren is not None:
                    mf = mode_fill(ren, box)
                    if mf:
                        fl = {"color": mf, "src": "sampled_unresolved"}

            ln = ln_props(el, scheme)
            if ln.get("src") == "default" and not has_theme_line(el):
                # ★ 没有 <a:ln>、主题也没给 lnRef → **原件没有边框**。
                #   以前这里会继续往下走去采样，把文字/填充当成了边框色。
                ln = dict(ln, src="noborder")
            if ln.get("src") == "scheme_unresolved":
                color_warnings.append("%s: 描边的主题色 %s 在原件主题里找不到"
                                      % (oid, ln.get("scheme")))
            if ln.get("color") is None and ln.get("src") in ("default", "unset",
                                                             "scheme_unresolved") \
                    and ren is not None:
                cands = []
                for k in range(1, 10):
                    cx = box[0] + box[2] * k / 10
                    cy = box[1] + box[3] * k / 10
                    c = darkest(ren, cx - 3, cy - 3, cx + 4, cy + 4)
                    if c:
                        cands.append(c)
                if cands:
                    ln = dict(ln, color=min(cands, key=lum), src="sampled")

            fg = {"color": None, "src": None}
            if base["text"]:
                fg = text_color_of(el, scheme)              # ★ XML 优先
                if fg.get("unknown_mods"):
                    color_warnings.append("%s: 文字色的修饰符 %s 没算，可能不准"
                                          % (oid, ",".join(fg["unknown_mods"])))
                if fg["src"] == "scheme_unresolved":
                    # ★ 原来这里**一声不响**：color 是 None，采样条件又不含
                    #   scheme_unresolved，于是下游静默把它变成 333333。
                    color_warnings.append(
                        "%s: 文字色的主题色 %s 在原件主题里找不到" % (oid, fg.get("scheme")))
                # 只有 **XML 里压根没写颜色** 才去渲染图采
                if fg["color"] is None and fg["src"] in ("unset",
                                                         "scheme_unresolved") \
                        and ren is not None:
                    sampled = text_color_sample(ren, box, fl.get("color"))
                    if sampled:
                        fg = {"color": sampled, "src": "sampled"}

            shape_kind = "text" if str(sh.shape_type).startswith("TEXT_BOX") else "shape"
            if base["prst"] in BRACE_PRSTS:
                shape_kind = "brace"
            elif base["prst"] in ARROW_PRSTS:
                shape_kind = "arrow"

            shapes.append(dict(base, kind=shape_kind,
                               fill=fl, line=ln,
                               font=font_of(el, scheme), insets=insets_of(el),
                               para=para_of(el), fg=fg))

    walk(slide.shapes, GroupCtx((0, 0), (1.0, 1.0), (0, 0)), [])

    textless = [s for s in shapes
                if s["kind"] in ("shape", "arrow") and not s["text"]]
    bound = [c for c in connectors if c["st_cxn"] and c["end_cxn"]]

    obs = {
        "observation_version": OBS_VERSION,
        "source": {
            "path": args.pptx, "slide": args.slide,
            "canvas": [round(v, 1) for v in canvas], "unit": "px@150dpi",
            "region": region, "render": args.render,
        },
        "shapes": shapes,
        "connectors": connectors,
        "groups": groups,
        "unsupported": unsupported,
        "color_warnings": color_warnings,
        "degrade": degrade,
        "stats": {
            "shapes": len(shapes),
            "connectors": len(connectors),
            "connectors_bound": len(bound),
            "groups": len(groups),
            "textless": len(textless),
            "unsupported": len(unsupported),
        },
    }
    with open(args.out, "w", encoding="utf-8") as f:
        json.dump(obs, f, ensure_ascii=False, indent=1)

    # ---------------- 人读的摘要 ----------------
    st = obs["stats"]
    print("观察到（第 %d 页）：" % args.slide)
    print("  形状 %d   连接符 %d（其中**端点绑定的 %d**）   组合 %d"
          % (st["shapes"], st["connectors"], st["connectors_bound"], st["groups"]))
    print("  没有文字的形状 %d" % st["textless"])
    print("  画不出来的 %d" % st["unsupported"])
    print()
    k = Counter(s["kind"] for s in shapes)
    print("  按类别:", dict(k))
    if connectors:
        print("  连接符端点情况：")
        print("     两端都绑定的  %d   ← 关系在文件里是明写的" % len(bound))
        print("     只绑一端的    %d"
              % len([c for c in connectors if bool(c["st_cxn"]) ^ bool(c["end_cxn"])]))
        print("     两端都没绑的  %d   ← ⚠ 方向和端点只能推，推出来必标「我推的」"
              % len([c for c in connectors if not c["st_cxn"] and not c["end_cxn"]]))
    print()
    print("  ⚠ 这只是**观察**，不是清单。")
    print("     「有 %d 个形状」和「图上有几个东西」**不是一回事**——" % st["shapes"])
    print("     归并成意义单位、判断关系，是下一步（要判断，可能要问用户）。")
    if textless:
        print()
        print("  ⚠ %d 个形状没有文字：先按「原件故意留空」处理，**不要自己填名字**，"
              % len(textless))
        print("     除非用户确认那是没画完。见 references/reading-rules.md")
    if degrade:
        print()
        print("  ⚠ 画不了、也没假装画的 %d 处：" % len(degrade))
        for d in degrade[:5]:
            print("     - %s" % d)
    if color_warnings:
        print()
        print("  ⚠ 搬运时主题色没解析干净的 %d 处（**没静默放过**）："
              % len(color_warnings))
        for w in color_warnings[:5]:
            print("     - %s" % w)
    if unsupported:
        print()
        print("  画不出来的：")
        for u in unsupported[:6]:
            print("     - %s  (%s)" % (u["what"], u["name"]))
        if len(unsupported) > 6:
            print("     … 还有 %d 个" % (len(unsupported) - 6))
    print()
    print("->", args.out)


if __name__ == "__main__":
    main()
