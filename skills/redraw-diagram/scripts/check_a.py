#!/usr/bin/env python3
"""
check_a.py —— Check A：**输出 vs 清单**，逐个对象对账。

它查的是「**执行有没有走样**」：清单说画什么，画出来了没有、画对没有。
它**不查**「清单对不对」——那是 Check B（清单 vs 原件）的事。

★ 为什么能机器做：
  生成时把清单的 id 写进了形状名字（见 build_output.py 的 stamp），
  所以两边能**逐个对象**对上，而不是只能比"文字集合"。
  这条就是"生成时留痕"。

对什么账（每一项都对着错误类型表）：
  对象数     清单 layout 的条目  vs  输出里的形状
  文字       清单 text.lines      vs  输出里该对象上的文字
  位置       清单 box / p1,p2     vs  输出里该对象的框
  预设几何   清单 prst            vs  输出里的 prstGeom

用法：
    python3 check_a.py work/spec.json work/out.pptx
    python3 check_a.py work/spec.json work/out.pptx -v
"""

import argparse
import re
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu

A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
DEFAULT_DPI = 150.0
EMU_PER_INCH = 914400
BOX_TOL_PX = 1.5      # 取整误差，1px @150dpi 之内算一致


def load(path):
    import json
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def norm(s):
    """规范化：去掉所有空白。比的是**字符序列和顺序**，不是分段方式。"""
    return re.sub(r"\s+", "", str(s or ""))


def emu_to_px(v, dpi=DEFAULT_DPI):
    return float(v) / EMU_PER_INCH * dpi


def shape_text(sh):
    try:
        return sh.text_frame.text if sh.has_text_frame else ""
    except Exception:
        return ""


def prst_of(sh):
    for e in sh._element.iter():
        if e.tag == A + 'prstGeom':
            return e.get('prst') or ''
    return ''


def is_line(sh):
    return (sh._element.tag.endswith('}cxnSp')
            or str(sh.shape_type).startswith("LINE"))


def line_pts(sh):
    x, y = emu_to_px(sh.left), emu_to_px(sh.top)
    w, h = emu_to_px(sh.width), emu_to_px(sh.height)
    x1, y1, x2, y2 = x, y, x + w, y + h
    el = sh._element
    for e in el.iter():
        if e.tag == A + 'xfrm':
            if e.get('flipH') in ('1', 'true'):
                x1, x2 = x2, x1
            if e.get('flipV') in ('1', 'true'):
                y1, y2 = y2, y1
            break
    return [x1, y1], [x2, y2]


def close(a, b, tol=BOX_TOL_PX):
    return abs(a - b) <= tol


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("spec")
    ap.add_argument("pptx")
    ap.add_argument("-v", "--verbose", action="store_true")
    args = ap.parse_args()

    spec = load(args.spec)
    prs = Presentation(args.pptx)
    slide = prs.slides[0]

    # 输出侧：按名字（= 清单 id）索引
    got = {}
    unnamed = []
    for sh in slide.shapes:
        nm = (sh.name or "").strip()
        if nm:
            got.setdefault(nm, []).append(sh)
        else:
            unnamed.append(sh)

    problems = []
    checked = 0

    for it in spec.get("layout") or []:
        sid = str(it["id"])
        cand = got.get(sid)
        if not cand:
            problems.append((sid, "没找到这个对象（清单里有，输出里没有）"))
            continue
        if len(cand) > 1:
            problems.append((sid, "输出里有 %d 个同名对象" % len(cand)))
        sh = cand[0]
        checked += 1

        # ---- 文字 ----
        want = norm("".join(str(x) for x in (it.get("text") or {}).get("lines") or []))
        have = norm(shape_text(sh))
        if want != have:
            problems.append((sid, "文字不一致：清单 %r  输出 %r" % (want, have)))

        # ---- 位置 ----
        if it["kind"] == "line":
            wp1, wp2 = it["p1"], it["p2"]
            gp1, gp2 = line_pts(sh)
            if not all(close(a, b) for a, b in zip(wp1 + wp2, gp1 + gp2)):
                problems.append((sid, "线的端点不一致：清单 %s→%s  输出 %s→%s"
                                 % (wp1, wp2, [round(v) for v in gp1],
                                    [round(v) for v in gp2])))
        elif it.get("box"):
            wb = it["box"]
            gb = [emu_to_px(sh.left), emu_to_px(sh.top),
                  emu_to_px(sh.width), emu_to_px(sh.height)]
            if not all(close(a, b) for a, b in zip(wb, gb)):
                problems.append((sid, "位置/尺寸不一致：清单 %s  输出 %s"
                                 % ([round(v, 1) for v in wb],
                                    [round(v, 1) for v in gb])))

        # ---- 预设几何 ----
        if it.get("prst") and it["kind"] not in ("line", "picture", "text"):
            gp = prst_of(sh)
            if gp and gp != it["prst"]:
                problems.append((sid, "预设几何不一致：清单 %r  输出 %r"
                                 % (it["prst"], gp)))

    # ---- 反向：输出里有、清单里没有 ----
    spec_ids = {str(x["id"]) for x in spec.get("layout") or []}
    extra = [nm for nm in got if nm not in spec_ids]
    if extra:
        problems.append(("(输出多出来的)", "清单里没有这些 id：%s" % ", ".join(extra)))
    if unnamed:
        problems.append(("(输出里没名字的)", "%d 个对象没有名字，无法对账" % len(unnamed)))

    # ---- 汇总 ----
    total = len(spec.get("layout") or [])
    print("Check A：输出 vs 清单")
    print("  清单 layout 条目   %d" % total)
    print("  输出里的形状       %d" % len(slide.shapes))
    print("  对上账的           %d" % checked)
    print()
    if problems:
        print("✗ 不一致 %d 处：" % len(problems))
        for sid, msg in problems:
            print("   %-14s %s" % (sid, msg))
        return 1
    print("✓ 全部对得上——执行没有走样")
    if args.verbose:
        print()
        print("  （这**不代表**清单是对的。清单对不对要靠 Check B：回原件对账。）")
    return 0


if __name__ == "__main__":
    sys.exit(main())
