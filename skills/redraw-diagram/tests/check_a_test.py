#!/usr/bin/env python3
"""
check_a_test.py —— 证明 check_a.py **有牙齿**。

只看"干净产物通过"是不够的：一个永远返回"通过"的检查也能做到。
所以拿生成的 pptx 手工破坏，看它抓不抓得到。

这比"再生成一次"更有说服力：破坏发生在**输出侧**，
正是 Check A 声称能看见的东西。

用法：
    python3 tests/check_a_test.py
"""

import copy
import os
import shutil
import subprocess
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
SCRIPTS = os.path.join(os.path.dirname(HERE), "scripts")
SPEC = os.path.join(HERE, "fixture-min.json")

from pptx import Presentation          # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.util import Emu, Inches       # noqa: E402

A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
PY = sys.executable


def run(*cmd):
    r = subprocess.run(cmd, capture_output=True, text=True)
    return r.returncode, r.stdout + r.stderr


def shapes_by_name(prs):
    return {s.name: s for s in prs.slides[0].shapes}


# --------------------------------------------------------------------------
MUTATIONS = [
    ("改一个形状里的一个字",
     lambda prs: _set_text(shapes_by_name(prs)["s02"], "制品库X")),

    ("把一个形状挪走 20px",
     lambda prs: _move(shapes_by_name(prs)["s01"], 20, 0)),

    ("删掉一个形状",
     lambda prs: _delete(shapes_by_name(prs)["s03"])),

    ("多画一个清单里没有的形状",
     lambda prs: _add_extra(prs)),

    ("改掉一个形状的预设几何",
     lambda prs: _set_prst(shapes_by_name(prs)["s01"], "ellipse")),

    ("把两个形状的文字互换",
     lambda prs: _swap_text(shapes_by_name(prs)["s01"],
                            shapes_by_name(prs)["s02"])),

    ("把线的端点反过来",
     lambda prs: _flip_line(shapes_by_name(prs)["s04"])),

    ("改掉一个对象的填充色",
     lambda prs: _set_fill(shapes_by_name(prs)["s01"], "FF0000")),
    ("把对象的填充整个去掉",
     lambda prs: _clear_fill(shapes_by_name(prs)["s01"])),
]


def _set_text(sh, s):
    sh.text_frame.paragraphs[0].runs[0].text = s


def _move(sh, dx, dy):
    sh.left = Emu(int(sh.left + Inches(dx / 150.0)))
    sh.top = Emu(int(sh.top + Inches(dy / 150.0)))


def _delete(sh):
    sh._element.getparent().remove(sh._element)


def _add_extra(prs):
    prs.slides[0].shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1),
                                   Inches(1), Inches(0.5)).name = "s99"


def _set_prst(sh, name):
    for e in sh._element.iter():
        if e.tag == A + 'prstGeom':
            e.set('prst', name)
            return


def _swap_text(a, b):
    ta = a.text_frame.text
    tb = b.text_frame.text
    _set_text(a, tb)
    _set_text(b, ta)


def _set_fill(sh, hexv):
    from pptx.dml.color import RGBColor
    sh.fill.solid()
    sh.fill.fore_color.rgb = RGBColor.from_string(hexv)


def _clear_fill(sh):
    sh.fill.background()


def _flip_line(sh):
    for e in sh._element.iter():
        if e.tag == A + 'xfrm':
            e.set('flipH', '1')
            return


# --------------------------------------------------------------------------
def main():
    tmp = tempfile.mkdtemp(prefix="checkA-")
    try:
        out = os.path.join(tmp, "out.pptx")
        rc, msg = run(PY, os.path.join(SCRIPTS, "build_output.py"), SPEC, "--pptx", out)
        if rc != 0:
            print("生成失败：", msg)
            return 1
        rc, msg = run(PY, os.path.join(SCRIPTS, "check_a.py"), SPEC, out)
        if rc != 0:
            print("✗ 干净产物就通不过，变异测试没意义：")
            print(msg)
            return 1
        print("基线：干净产物通过 ✓")

        caught, escaped = 0, []
        for name, fn in MUTATIONS:
            dst = os.path.join(tmp, "m.pptx")
            shutil.copyfile(out, dst)
            prs = Presentation(dst)
            try:
                fn(prs)
                prs.save(dst)
            except Exception as e:
                escaped.append((name, "变异执行失败：%r" % e))
                continue
            rc, msg = run(PY, os.path.join(SCRIPTS, "check_a.py"), SPEC, dst)
            if rc != 0:
                caught += 1
                detail = [l.strip() for l in msg.splitlines() if l.strip().startswith(tuple("s0123456789"))]
                print("  ✓ %-24s → %s" % (name, (detail[0] if detail else msg.strip()[:70])))
            else:
                escaped.append((name, "**Check A 没抓到**"))

        print()
        print("变异测试：%d/%d 被抓到" % (caught, len(MUTATIONS)))
        if escaped:
            print()
            print("✗ 漏掉的：")
            for name, why in escaped:
                print("   %-24s %s" % (name, why))
            return 1
        print("✓ 全部抓到")
        return 0
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


if __name__ == "__main__":
    sys.exit(main())
