#!/usr/bin/env python3
"""
check_b_test.py —— 证明 check_b.py **有牙齿**。

自包含：临时造一份小 pptx 当"原件"，再配一份和它对得上的清单。
基线必须通过，然后从**两边**下手破坏，看它抓不抓得到。

从"原件侧"破坏很重要 —— 那才是 Check B 和 Check A 的分水岭：
Check A 对原件的改动完全无感。

用法：
    python3 tests/check_b_test.py
"""

import copy
import json
import os
import shutil
import subprocess
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
SCRIPTS = os.path.join(os.path.dirname(HERE), "scripts")
PY = sys.executable

from pptx import Presentation          # noqa: E402
from pptx.util import Inches           # noqa: E402

TEXTS = ["开发环境", "制品库", "预发布环境"]


def run(*cmd):
    r = subprocess.run(cmd, capture_output=True, text=True)
    return r.returncode, r.stdout + r.stderr


def make_pptx(path, texts=TEXTS):
    prs = Presentation()
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    for i, t in enumerate(texts):
        sl.shapes.add_textbox(Inches(0.5 + i * 2), Inches(1), Inches(1.6), Inches(0.6)) \
            .text_frame.text = t
    prs.save(path)


def make_spec(path, texts=TEXTS, prov="original"):
    spec = {
        "spec_version": "0.2",
        "source": {"canvas": [2000, 1125], "unit": "px@150dpi", "files": []},
        "entities": [
            {"id": "e%02d" % i, "role": ["shape"], "text": t,
             "provenance": prov, "evidence": ["p:sp#%d" % i], "status": "confirmed"}
            for i, t in enumerate(texts)
        ],
        "relations": [], "groups": [], "layout": [],
        "changes": [], "open_questions": [], "findings": [], "unsupported": [],
    }
    with open(path, "w", encoding="utf-8") as f:
        json.dump(spec, f, ensure_ascii=False)
    return spec


# --------------------------------------------------------------------------
# 每个变异 = (名字, 改清单的函数 或 None, 改原件的函数 或 None)
MUTATIONS = [
    ("改清单里一个字",
     lambda s: s["entities"][0].__setitem__("text", "开发环坑"), None),
    ("清单里加一条假的原件文字",
     lambda s: s["entities"].append(
         {"id": "e99", "role": ["shape"], "text": "生产环境",
          "provenance": "original", "evidence": ["p:sp#99"], "status": "confirmed"}), None),
    ("清单里把一条文字删空",
     lambda s: s["entities"][1].__setitem__("text", None), None),
    ("清单里把一条文字截成半截",
     lambda s: s["entities"][2].__setitem__("text", "预发布"), None),
    ("原件里改掉一个字（清单没动）",
     None, lambda texts: texts.__setitem__(0, "开发环坑")),
    ("原件里多一个文本框（清单没动）",
     None, lambda texts: texts.append("生产环境")),
    ("原件里删掉一个文本框（清单没动）",
     None, lambda texts: texts.pop(0)),
]


def should_not_trigger():
    """这条**不应该**报 —— 用来证明过滤器是对的，不是永远报警。"""
    return ("标成「用户说的」就不查（设计如此）",
            lambda s: s["entities"][0].__setitem__("provenance", "user"), None)


def main():
    tmp = tempfile.mkdtemp(prefix="checkB-")
    try:
        pptx = os.path.join(tmp, "orig.pptx")
        spec = os.path.join(tmp, "spec.json")
        make_pptx(pptx)
        make_spec(spec)

        rc, msg = run(PY, os.path.join(SCRIPTS, "check_b.py"), spec, pptx)
        if rc != 0 or "对得上" not in msg:
            print("✗ 基线就不通过，变异测试没意义：")
            print(msg)
            return 1
        print("基线：对得上的清单通过 ✓")

        caught, escaped = 0, []
        for name, mut_spec, mut_orig in MUTATIONS:
            sp = os.path.join(tmp, "m.json")
            px = os.path.join(tmp, "m.pptx")
            s = make_spec(sp)
            texts = list(TEXTS)
            if mut_spec:
                mut_spec(s)
                with open(sp, "w", encoding="utf-8") as f:
                    json.dump(s, f, ensure_ascii=False)
            if mut_orig:
                mut_orig(texts)
            make_pptx(px, texts)

            rc, msg = run(PY, os.path.join(SCRIPTS, "check_b.py"), sp, px, "--strict")
            if rc != 0:
                caught += 1
                line = next((l.strip() for l in msg.splitlines()
                             if l.strip().startswith(("✗", "⚠"))), msg.strip()[:60])
                print("  ✓ %-26s → %s" % (name, line))
            else:
                escaped.append((name, "**Check B 没抓到**"))

        # 反向：标成「用户说的」不该报
        name, mut_spec, _ = should_not_trigger()
        sp = os.path.join(tmp, "n.json")
        s = make_spec(sp)
        mut_spec(s)
        with open(sp, "w", encoding="utf-8") as f:
            json.dump(s, f, ensure_ascii=False)
        rc, msg = run(PY, os.path.join(SCRIPTS, "check_b.py"), sp, pptx, "--strict")
        if rc == 0:
            print("  ✓ %-26s → 没报（正确）" % name)
        else:
            escaped.append((name, "**不该报却报了**"))

        print()
        total = len(MUTATIONS) + 1
        print("变异测试：%d/%d 符合预期" % (caught + (1 if rc == 0 else 0), total))
        if escaped:
            print()
            print("✗ 不符合预期的：")
            for n, why in escaped:
                print("   %-26s %s" % (n, why))
            return 1
        print("✓ 全部符合预期")
        return 0
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


if __name__ == "__main__":
    sys.exit(main())
