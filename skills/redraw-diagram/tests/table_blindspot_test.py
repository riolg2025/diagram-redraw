#!/usr/bin/env python3
"""
table_blindspot_test.py —— 证明「表格盲区」这个洞被堵上了。

## 这个洞是什么（实测撞出来的）

拿 deck 的第 1 页试「换一类图」，发现：

    原件第 1 页   32 个形状 = 17 自选图形 + **10 个表格** + 5 文本框
    观察工具     只报 22 个，而 stats 写着 **"unsupported": 0** ✗✗

**表格被静默吞了，它还说"没漏"** ✗ —— R3.5 那类病（「我不知道」和「没有」没分开）。

根因两层：
  ① `Xfrm` 只在 `a:` 命名空间里找 `xfrm` ✗，而表格（`p:graphicFrame`）
     的变换是 **`p:xfrm`** → 在第 583 行 `continue` 掉了
  ② 那个 `continue` **没有记账** ✗

另一头 **Check B 也是瞎的**：`if not sh.has_text_frame: continue` ✗
表格没有 `text_frame` ✗ → 原件那一侧**等于没有这些文字** → 它也不会报遗漏 ✗
**两头都瞎，所以这个洞以前从没被发现过** ✗

## 这个测试守什么

★ **核心不变式：账要对上** —— `观察到的形状 + 记下来的 = 原件实际数量` ✓
   这条比"支持表格"更重要：**支持不了不要紧，不许说没有** ✓

用法：
    python3 tests/table_blindspot_test.py
"""

import json
import os
import subprocess
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
SCRIPTS = os.path.join(os.path.dirname(HERE), "scripts")
PY = sys.executable
sys.path.insert(0, SCRIPTS)

from pptx import Presentation            # noqa: E402
from pptx.util import Emu, Inches        # noqa: E402


def build_deck(path):
    """造一页：1 个普通形状 + 1 个 2×3 表格 + 1 个文本框。"""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    s = prs.slides.add_slide(prs.slide_layouts[6])

    tb = s.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(0.6))
    tb.text_frame.text = "标题文字"

    sh = s.shapes.add_shape(1, Inches(1), Inches(2), Inches(2), Inches(0.8))
    sh.text_frame.text = "一个形状"

    gf = s.shapes.add_table(2, 3, Inches(4), Inches(2), Inches(5), Inches(1.2))
    t = gf.table
    for i, v in enumerate(["甲", "乙", "丙", "丁", "戊", "己"]):
        t.cell(i // 3, i % 3).text = v

    prs.save(path)
    return 3          # 真实对象数


def main():
    ok = True
    tmp = tempfile.mkdtemp()
    deck = os.path.join(tmp, "t.pptx")
    real = build_deck(deck)

    print("表格盲区测试（造一页：1 形状 + 1 表格 2×3 + 1 文本框 = %d 个对象）\n" % real)

    # ── ① 观察工具：不许静默吞 ──────────────────────────────────────
    out = os.path.join(tmp, "obs.json")
    r = subprocess.run([PY, os.path.join(SCRIPTS, "observe_pptx_shapes.py"),
                        deck, "--slide", "1", "--out", out],
                       capture_output=True, text=True)
    if r.returncode != 0:
        print("  ✗ 观察工具跑失败：%s" % r.stderr[-300:])
        return 1
    O = json.load(open(out, encoding="utf-8"))
    st = O.get("stats") or {}
    cells = [x for x in O.get("shapes") or [] if x.get("table")]
    tables = [g for g in O.get("groups") or [] if g.get("kind") == "table"]
    # ★ 表格**展开成格子**了，所以"账"不能按 shapes 直接数 ——
    #   一个表格算**一个**原件对象，它的格子是**内容**不是新对象。
    seen = (st.get("shapes", 0) - len(cells)) + len(tables) + st.get("unsupported", 0)
    print("  ① 每个原件对象都被交代了吗")
    print("     非表格形状 %d + 表格 %d + 记下来的 %d = %d，原件 %d  %s"
          % (st.get("shapes", 0) - len(cells), len(tables), st.get("unsupported", 0),
             seen, real, "✓" if seen == real else "✗ **有对象无声无息地没了**"))
    if seen != real:
        ok = False
    # ★ 表格现在**读得出来了** —— 所以它既不该在 unsupported 里，
    #   也不该被静默吞掉：它必须**变成格子**出现在 shapes 里。
    if cells:
        print("     ✓ 表格读成了 %d 个格子（不是记成「不支持」，更不是吞掉）" % len(cells))
    elif any("表格" in str(u.get("what", "")) for u in O.get("unsupported") or []):
        print("     ⚠ 表格仍被记成「不支持」——账是对的，但能力还没接上")
    else:
        print("     ✗ 表格既没读出来、也没记下来（**又被静默吞了**）")
        ok = False

    # ── ② Check B：原件那侧要读得到表格文字 ────────────────────────
    print()
    print("  ② Check B 读得到表格文字吗")
    import importlib.util
    spec = importlib.util.spec_from_file_location(
        "cb", os.path.join(SCRIPTS, "check_b.py"))
    cb = importlib.util.module_from_spec(spec)
    sys.modules["cb"] = cb
    spec.loader.exec_module(cb)
    blocks = cb.extract_pptx_text(deck, 1)
    blob = "".join(blocks)
    miss = [c for c in "甲乙丙丁戊己" if c not in blob]
    print("     取到 %d 块文字" % len(blocks))
    if miss:
        print("     ✗ 表格里的字没取到：%s" % "、".join(miss))
        ok = False
    else:
        print("     ✓ 表格里 6 个单元格的字全取到了（这样它才报得出「遗漏」）")

    print()
    print("table_blindspot_test：%s" % ("全部符合预期 ✓" if ok else "有失败 ✗"))
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
