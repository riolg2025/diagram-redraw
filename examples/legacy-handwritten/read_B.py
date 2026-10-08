#!/usr/bin/env python3
"""
dryrun-02：把 B 区（组合 130）完整读出来。

关键：组内元素给的是**组内相对坐标**，必须做变换才能落到页面上。
  abs = off + (child - chOff) * (ext / chExt)

输出 inventory-B.json + 控制台表格。
"""
import json
from pptx import Presentation
from lxml import etree

P = '{http://schemas.openxmlformats.org/presentationml/2006/main}'
A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
EMU_PX = 6096.0


def px(v):
    return round(v / EMU_PX, 1)


def get_xfrm(el):
    for e in el.iter():
        if e.tag == A + 'xfrm':
            o, ex = e.find(A + 'off'), e.find(A + 'ext')
            co, ce = e.find(A + 'chOff'), e.find(A + 'chExt')
            if o is None or ex is None:
                continue
            d = {'off': (int(o.get('x')), int(o.get('y'))),
                 'ext': (int(ex.get('cx')), int(ex.get('cy')))}
            if co is not None and ce is not None:
                d['chOff'] = (int(co.get('x')), int(co.get('y')))
                d['chExt'] = (int(ce.get('cx')), int(ce.get('cy')))
            return d
    return None


def win_xfrm(sh):
    """shape 自己的 a:xfrm（spPr 里的那个）"""
    spPr = None
    for e in sh._element.iter():
        if e.tag == P + 'spPr' or e.tag == A + 'spPr':
            spPr = e
            break
    if spPr is None:
        return None
    x = spPr.find(A + 'xfrm')
    if x is None:
        return None
    o, ex = x.find(A + 'off'), x.find(A + 'ext')
    if o is None or ex is None:
        return None
    return (int(o.get('x')), int(o.get('y')), int(ex.get('cx')), int(ex.get('cy')))


def n_runs(sh):
    tb = sh._element.find(P + 'txBody')
    if tb is None:
        return None
    return len(tb.findall('.//' + A + 'r'))


def prst(sh):
    for e in sh._element.iter():
        if e.tag.endswith('}prstGeom'):
            return e.get('prst') or ''
    return ''


def main():
    pr = Presentation("/Users/rio/Desktop/测试架构图.pptx")
    grp = [sh for sh in pr.slides[4].shapes if sh.shape_type == 6][1]
    gx = get_xfrm(grp._element)
    print("组变换:", {k: (px(v[0]), px(v[1])) for k, v in gx.items() if k in ('off', 'ext')})
    if 'chOff' in gx:
        print("         chOff/chExt:", gx['chOff'], gx['chExt'])
        sx = gx['ext'][0] / gx['chExt'][0]
        sy = gx['ext'][1] / gx['chExt'][1]
    else:
        sx = sy = 1.0
    print(f"         缩放 sx={sx:.4f} sy={sy:.4f}")

    rows = []

    def walk(shapes, base_off, sxx, syy, ch_off):
        for sh in shapes:
            w = win_xfrm(sh)
            if w is None:
                continue
            # 组内相对坐标 -> 页面坐标：必须减去 chOff，否则整体平移
            ax = base_off[0] + (w[0] - ch_off[0]) * sxx
            ay = base_off[1] + (w[1] - ch_off[1]) * syy
            aw = w[2] * sxx
            ah = w[3] * syy
            st = str(sh.shape_type).split(' ')[0]
            r = n_runs(sh)
            t = sh.text_frame.text.strip().replace("\n", " / ") if sh.has_text_frame else ""
            rows.append({
                "name": sh.name, "type": st, "prst": prst(sh),
                "runs": r, "text": t,
                "box_px": [px(ax), px(ay), px(ax + aw), px(ay + ah)],
                "w_px": px(aw), "h_px": px(ah),
            })
            if st == "GROUP":
                g2 = get_xfrm(sh._element)
                if g2 and 'chOff' in g2:
                    s2x = g2['ext'][0] / g2['chExt'][0]
                    s2y = g2['ext'][1] / g2['chExt'][1]
                    walk(sh.shapes,
                         (base_off[0] + (g2['off'][0] - ch_off[0]) * sxx,
                          base_off[1] + (g2['off'][1] - ch_off[1]) * syy),
                         s2x, s2y, g2['chOff'])
                else:
                    walk(sh.shapes, (ax, ay), sxx, syy, (0, 0))

    walk(grp.shapes, gx['off'], sx, sy, gx.get('chOff', (0, 0)))

    json.dump(rows, open("diagram-redraw/dryrun-02/inventory-B.json", "w"),
              ensure_ascii=False, indent=1)

    rows.sort(key=lambda r: (round(r["box_px"][1] / 25), r["box_px"][0]))
    print(f"\n共 {len(rows)} 个对象（含组内）\n")
    print(f"{'类型':11s} {'几何':15s} {'x':>6s}{'y':>6s}{'w':>6s}{'h':>6s}  {'字':4s} 内容")
    for r in rows:
        flag = "空" if r["runs"] == 0 else (str(r["runs"]) if r["runs"] else "-")
        print(f'{r["type"]:11s} {r["prst"][:14]:15s} '
              f'{r["box_px"][0]:6.0f}{r["box_px"][1]:6.0f}{r["w_px"]:6.0f}{r["h_px"]:6.0f}  '
              f'{flag:4s} {r["text"][:44]}')


if __name__ == "__main__":
    main()
