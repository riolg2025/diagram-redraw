#!/usr/bin/env python3
"""
dryrun-01 第一步（v2）：把第 6 页读到的东西数清楚，标在原图上。

v1 -> v2 改了两处，都是 v1 自己暴露出来的：
  1. 连线编号原来放在 bbox 中点，结果压在文字上、而且分不清对哪条线。
     v2 改成：按 flipH/flipV 算出线的真实两端，取线段中点，垂直偏移，并拉一条引线。
  2. 分区编号原来放在框内左上角，压在框里的内容上。v2 移到了框外。

产物：inventory.json / slide6-marked.png
"""

import json
from PIL import Image, ImageDraw, ImageFont

EMU_PER_PX = 6096.0
SRC = "diagram-redraw/dryrun-01/inventory-raw.json"
PNG = "diagram-redraw/dryrun-01/pages/page-0001.png"
OUT_PNG = "diagram-redraw/dryrun-01/slide6-marked.png"
OUT_JSON = "diagram-redraw/dryrun-01/inventory.json"

CONTAINERS = {"矩形 63", "矩形 72", "矩形 12", "矩形 15", "矩形 121", "Rectangle 8"}
TITLE = {"文本占位符 2"}


def flips(sh):
    """读翻转标记；拿不到就当没有。"""
    try:
        xfrm = None
        for child in sh._element.iter():
            if child.tag.endswith("}xfrm"):
                xfrm = child
                break
        if xfrm is None:
            return False, False
        return (xfrm.get("flipH") in ("1", "true"),
                xfrm.get("flipV") in ("1", "true"))
    except Exception:
        return False, False


def classify(r):
    t, name, text = r["type"], r["name"], r["text"]
    if r["depth"] > 0:
        return "图标内部"
    if t.startswith("LINE"):
        return "连线"
    if t.startswith("PICTURE"):
        return "截图"
    if t.startswith("GROUP"):
        return "节点"
    if name in TITLE:
        return "标题"
    if name in CONTAINERS:
        return "分区"
    if t.startswith("AUTO_SHAPE"):
        return "节点" if text else "装饰"
    if t.startswith("TEXT_BOX"):
        return "文字"
    return "其他"


def main():
    recs = json.load(open(SRC))
    for r in recs:
        r["cls"] = classify(r)
    top = [r for r in recs if r["depth"] == 0]

    im = Image.open(PNG).convert("RGB")
    d = ImageDraw.Draw(im)
    f = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial Bold.ttf", 19)

    def dot(cx, cy, label, fill, leader_to=None):
        if leader_to:
            d.line([cx, cy, leader_to[0], leader_to[1]], fill=fill, width=3)
        r = 14
        d.ellipse([cx - r, cy - r, cx + r, cy + r], fill=fill)
        d.ellipse([cx - r - 1, cy - r - 1, cx + r + 1, cy + r + 1],
                  outline=(30, 30, 30), width=2)
        w = d.textlength(label, font=f)
        d.text((cx - w / 2, cy - 11), label, fill=(255, 255, 255), font=f)

    inv = {"nodes": [], "connectors": [], "containers": [],
           "pictures": [], "texts": [], "title": [], "skipped": []}

    n = g = 0
    for r in sorted(top, key=lambda r: (round(r["top"] / 30), r["left"])):
        x, y, w, h = r["left"], r["top"], r["w"], r["h"]
        if r["cls"] == "节点":
            n += 1
            lbl = f"N{n}"
            dot(x + 9, y + 9, lbl, (206, 38, 38))
            inv["nodes"].append({"id": lbl, "name": r["name"],
                                 "text": r["text"][:40],
                                 "box_px": [x, y, x + w, y + h]})
        elif r["cls"] == "分区":
            g += 1
            lbl = f"G{g}"
            d.rectangle([x, y, x + w, y + h], outline=(21, 94, 190), width=4)
            ty = y - 20 if y - 20 > 20 else y + h + 20
            dot(x + 6, ty, lbl, (21, 94, 190), leader_to=(x + 2, y))
            inv["containers"].append({"id": lbl, "name": r["name"],
                                      "box_px": [x, y, x + w, y + h]})

    from pptx import Presentation
    pr = Presentation("/Users/rio/Desktop/测试架构图.pptx")
    lines = [sh for sh in pr.slides[5].shapes
             if str(sh.shape_type).startswith("LINE")]
    c = 0
    for sh in lines:
        c += 1
        lbl = f"C{c}"
        x = sh.left / EMU_PER_PX
        y = sh.top / EMU_PER_PX
        w = sh.width / EMU_PER_PX
        h = sh.height / EMU_PER_PX
        fh, fv = flips(sh)
        x1, y1, x2, y2 = x, y, x + w, y + h
        if fh:
            x1, x2 = x2, x1
        if fv:
            y1, y2 = y2, y1
        mx, my = (x1 + x2) / 2, (y1 + y2) / 2
        dx, dy = x2 - x1, y2 - y1
        L = max((dx * dx + dy * dy) ** 0.5, 1)
        nx, ny = -dy / L, dx / L
        OFF = 24
        lx = min(max(mx + nx * OFF, 18), 1982)
        ly = min(max(my + ny * OFF, 18), 1107)
        dot(lx, ly, lbl, (222, 122, 12), leader_to=(mx, my))
        inv["connectors"].append({"id": lbl, "name": sh.name,
                                  "endpoints_px": [round(x1), round(y1),
                                                   round(x2), round(y2)]})

    for r in top:
        x, y, w, h = r["left"], r["top"], r["w"], r["h"]
        if r["cls"] == "截图":
            inv["pictures"].append({"name": r["name"],
                                    "box_px": [x, y, x + w, y + h]})
        elif r["cls"] == "文字" and r["text"]:
            inv["texts"].append({"name": r["name"], "text": r["text"][:60]})
        elif r["cls"] == "标题":
            inv["title"].append({"name": r["name"], "text": r["text"][:60]})
        elif r["cls"] == "装饰":
            inv["skipped"].append({"name": r["name"], "cls": r["cls"]})

    inv["counts"] = {"节点": n, "连线": c, "分区": g,
                     "截图": len(inv["pictures"]), "文字": len(inv["texts"])}
    json.dump(inv, open(OUT_JSON, "w"), ensure_ascii=False, indent=1)
    im.save(OUT_PNG)
    print("计数:", inv["counts"])


if __name__ == "__main__":
    main()
